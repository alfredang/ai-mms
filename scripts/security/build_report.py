#!/usr/bin/env python3
"""Render a security-scanner findings JSON file into a Word assessment report.

Usage: python3 scripts/security/build_report.py <findings.json> <out.docx>

The JSON is written by the `security-scanner` agent (.claude/agents/security-scanner.md);
its schema is documented there. Reports are confidential and stay local
(security-assessment/ is gitignored) — never upload them anywhere.
"""
import json
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

SEVERITIES = ["Critical", "High", "Medium", "Low", "Info"]
SEV_COLOR = {
    "Critical": "7B1E1E", "High": "C0392B", "Medium": "E67E22",
    "Low": "2E86C1", "Info": "7F8C8D",
}
NAVY = RGBColor(0x1F, 0x3A, 0x5F)


def shade(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def table(doc, header, rows, widths=None, sev_col=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(c, "1F3A5F")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
            if sev_col is not None and i == sev_col and v in SEV_COLOR:
                shade(cells[i], SEV_COLOR[v])
                for p in cells[i].paragraphs:
                    for r in p.runs:
                        r.bold = True
                        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
    return t


def field(doc, label, value):
    if not value:
        return
    p = doc.add_paragraph()
    p.add_run(f"{label}: ").bold = True
    p.add_run(str(value))


def main(src, out):
    with open(src, encoding="utf-8") as f:
        data = json.load(f)
    findings = sorted(
        data.get("findings", []),
        key=lambda x: (SEVERITIES.index(x.get("severity", "Info")), x.get("id", "")),
    )
    counts = {s: sum(1 for x in findings if x.get("severity") == s) for s in SEVERITIES}

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    # Cover
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("CONFIDENTIAL — INTERNAL USE ONLY")
    r.bold = True
    r.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
    title = doc.add_heading("Security Assessment Report", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(data.get("system", "Tertiary Infotech Academy — Magento Management System (AI-MMS)"))
    r.font.size = Pt(14)
    r.font.color.rgb = NAVY
    table(doc, ["Item", "Detail"], [
        ["Assessment date", data.get("scan_date", "")],
        ["Assessor", data.get("assessor", "security-scanner agent (Claude Code)")],
        ["Repository commit", data.get("commit", "")],
        ["Targets", ", ".join(data.get("targets", []))],
        ["Overall risk rating", data.get("overall_rating", "")],
    ])
    doc.add_page_break()

    doc.add_heading("1. Executive Summary", 1)
    doc.add_paragraph(data.get("executive_summary", ""))
    table(doc, ["Severity", "Count", "Fixed", "Open"], [
        [s, counts[s],
         sum(1 for x in findings if x.get("severity") == s and x.get("status", "").startswith("Fixed")),
         sum(1 for x in findings if x.get("severity") == s and not x.get("status", "").startswith("Fixed"))]
        for s in SEVERITIES
    ], sev_col=0)

    doc.add_heading("2. Scope and Methodology", 1)
    doc.add_paragraph(data.get("scope", ""))
    doc.add_paragraph(data.get("methodology", ""))
    doc.add_heading("Severity classification", 2)
    table(doc, ["Severity", "CVSS v3.1", "Meaning"], [
        ["Critical", "9.0–10.0", "Directly exploitable; full compromise, data breach or payment fraud likely."],
        ["High", "7.0–8.9", "Exploitable with little effort; significant data or integrity impact."],
        ["Medium", "4.0–6.9", "Exploitable under specific conditions or limited impact; fix in normal cycle."],
        ["Low", "0.1–3.9", "Hardening gap; minimal direct impact."],
        ["Info", "0.0", "Observation or good practice, no direct risk."],
    ], sev_col=0)

    doc.add_heading("3. Findings Summary", 1)
    table(doc, ["ID", "Severity", "Title", "Category", "Status"], [
        [x.get("id", ""), x.get("severity", ""), x.get("title", ""),
         x.get("category", ""), x.get("status", "Open")]
        for x in findings
    ], sev_col=1)

    doc.add_heading("4. Detailed Findings", 1)
    for x in findings:
        doc.add_heading(f"{x.get('id', '')} — {x.get('title', '')}", 2)
        p = doc.add_paragraph()
        r = p.add_run(f" {x.get('severity', '')} ")
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(SEV_COLOR.get(x.get("severity"), "7F8C8D"))
        if x.get("cvss"):
            p.add_run(f"   CVSS {x['cvss']}")
        field(doc, "Category (OWASP Top 10)", x.get("category"))
        field(doc, "STRIDE", x.get("stride"))
        field(doc, "MITRE ATT&CK", x.get("attack"))
        field(doc, "CIA impact", x.get("cia"))
        field(doc, "Location", x.get("location"))
        field(doc, "Description", x.get("description"))
        field(doc, "Evidence", x.get("evidence"))
        field(doc, "Impact", x.get("impact"))
        field(doc, "Recommendation", x.get("recommendation"))
        field(doc, "Status", x.get("status", "Open"))
        field(doc, "Remediation notes", x.get("fix_notes"))

    if data.get("positive_controls"):
        doc.add_heading("5. Controls Verified In Place", 1)
        for c in data["positive_controls"]:
            doc.add_paragraph(c, style="List Bullet")

    if data.get("next_steps"):
        doc.add_heading("6. Recommended Next Steps", 1)
        for c in data["next_steps"]:
            doc.add_paragraph(c, style="List Number")

    doc.save(out)
    print(f"wrote {out}: {len(findings)} findings "
          + ", ".join(f"{s}={counts[s]}" for s in SEVERITIES))


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
